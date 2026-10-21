"""Haftalık ilerleme raporlarını ders şablonuna (Takım Kaptanı Geri Bildirim Formu + rapor gövdesi) uygun .docx olarak üretir."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell

TEAM_NAME = "BankTech"
PROJECT_TITLE = "Banka Pazarlama Kampanyalarına Müşteri Yanıtının Tahmini"
PROJECT_LINKS = [
    "https://github.com/mustafaemre027/bank-predict",
    "https://drive.google.com/drive/folders/1PwRRUcQtGplImXdSqaZmNaIWLk_eboNp",
]
TITLE_LINES = [
    "T. C.",
    "FIRAT ÜNİVERSİTESİ",
    "MÜHENDİSLİK FAKÜLTESİ",
    "BİLGİSAYAR MÜHENDİSLİĞİ",
    "BİLGİSAYAR MÜHENDİSLİĞİ TASARIM DERSİ PROJESİ KAPSAMINDA",
    "TAKIM ÇALIŞMASI PROJESİ",
]
TEAM_PROCESS_CRITERIA = [
    "Ekip içi iletişim düzenliydi",
    "Görev dağılımı netti",
    "Tüm üyeler sürece aktif katıldı",
    "Toplantılar planlandığı şekilde yapıldı",
    "Zaman yönetimi yeterliydi",
    "Ekip içi anlaşmazlıklar uygun şekilde çözüldü",
]
CAPTAIN_ROLE = "Takım Kaptanı"
CHECKED = "√"
UNCHECKED = "☐"

FONT_NAME = "Arial"
BODY_FONT_SIZE = Pt(10)
TABLE_FONT_SIZE = Pt(9)
BORDER_COLOR = "70AD47"
HEADER_FILL = "F2F2F2"
CAPTION_COLOR = RGBColor(0x0B, 0x0B, 0x0B)
PAGE_MARGIN = Cm(2.5)
A4_WIDTH = Cm(21)
A4_HEIGHT = Cm(29.7)
CONTENT_WIDTH_CM = 16.0
TWIPS_PER_CM = 567
FIGURE_WIDTH = Cm(15)
BORDER_SIZE_EIGHTHS = "6"
# OOXML şemasında tblPr alt öğeleri sabit sıradadır; tblBorders bunlardan önce gelmelidir.
TABLE_BORDER_SUCCESSORS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


@dataclass(frozen=True)
class Member:
    student_id: str
    full_name: str
    role: str


TEAM = [
    Member("230260054", "Ömer NARÇİÇEĞİ", "Makine Öğrenmesi Geliştirici"),
    Member("230260173", "Mesut ASLAN", "Veri Analisti / Yazılım Geliştirici"),
    Member("230260143", "Mustafa Emre KAPLAN", "Yazılım Geliştirici / Test ve Raporlama"),
]


def find_member(student_id: str) -> Member:
    for member in TEAM:
        if member.student_id == student_id:
            return member
    raise ValueError(f"'{student_id}' numaralı üye ekipte yok. Geçerli numaralar: {[m.student_id for m in TEAM]}")


def report_file_name(week: int) -> str:
    """Teslim dosyası adı: öğrenci numaraları sıralı, ardından hafta (2. hafta teslimiyle aynı biçim)."""
    student_ids = "_".join(sorted(member.student_id for member in TEAM))
    return f"{student_ids}_hafta{week}.docx"


def _set_table_borders(table: Table, color: str = BORDER_COLOR) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), BORDER_SIZE_EIGHTHS)
        element.set(qn("w:color"), color)
        borders.append(element)
    table_properties = table._tbl.tblPr
    successor = next((child for tag in TABLE_BORDER_SUCCESSORS if (child := table_properties.find(qn(tag))) is not None), None)
    if successor is None:
        table_properties.append(borders)
    else:
        successor.addprevious(borders)


def _set_column_widths(table: Table, widths_cm: list[float]) -> None:
    # Word sütun genişliğini tablo ızgarasından (gridCol), bazı görüntüleyiciler hücreden okur; ikisi birlikte ayarlanır.
    # autofit=False, şemaya uygun sırada sabit tablo yerleşimi (tblLayout fixed) ekler.
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(int(sum(widths_cm) * TWIPS_PER_CM)))
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def _equal_widths(column_count: int) -> list[float]:
    return [CONTENT_WIDTH_CM / column_count] * column_count


def _shade_cell(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _write_cell(cell: _Cell, text: str, bold: bool = False, size: Pt = TABLE_FONT_SIZE, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = size


class WeeklyReport:
    """Bir haftalık raporu bölüm bölüm kurar; tablo ve şekil numaralarını otomatik verir."""

    def __init__(self, week: int, captain_id: str) -> None:
        self.week = week
        self.captain = find_member(captain_id)
        self.document = Document()
        self._table_number = 0
        self._figure_number = 0
        self._configure_document()

    def _configure_document(self) -> None:
        normal = self.document.styles["Normal"]
        normal.font.name = FONT_NAME
        normal.font.size = BODY_FONT_SIZE
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        for section in self.document.sections:
            section.page_width = A4_WIDTH
            section.page_height = A4_HEIGHT
            section.left_margin = section.right_margin = PAGE_MARGIN
            section.top_margin = section.bottom_margin = PAGE_MARGIN

    def _new_table(self, rows: int, cols: int, widths_cm: list[float] | None = None) -> Table:
        table = self.document.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(table)
        _set_column_widths(table, widths_cm or _equal_widths(cols))
        return table

    def _bold_line(self, text: str, size: Pt = BODY_FONT_SIZE, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = size

    def _checklist(self, criteria: list[tuple[str, bool]]) -> None:
        table = self._new_table(rows=len(criteria) + 1, cols=3, widths_cm=[14.4, 0.8, 0.8])
        for cell, header in zip(table.rows[0].cells, ["Kriter", CHECKED, "X"]):
            _write_cell(cell, header, bold=True)
            _shade_cell(cell, HEADER_FILL)
        for row, (criterion, is_met) in zip(table.rows[1:], criteria):
            _write_cell(row.cells[0], criterion)
            _write_cell(row.cells[1], CHECKED if is_met else UNCHECKED, align=WD_ALIGN_PARAGRAPH.CENTER)
            _write_cell(row.cells[2], UNCHECKED if is_met else CHECKED, align=WD_ALIGN_PARAGRAPH.CENTER)

    def add_captain_form(self, team_process: list[tuple[str, bool]], progress: list[tuple[str, bool]]) -> None:
        self._bold_line("TAKIM KAPTANI GERİ BİLDİRİM FORMU", size=Pt(11))
        self._bold_line("Bölüm 1 - Takım Bilgileri")
        members = [m for m in TEAM if m != self.captain]
        info = [
            ("Ekip Adı:", TEAM_NAME),
            ("Takım Kaptanı:", f"{self.captain.full_name} ({self.captain.student_id})"),
            ("Ekip Üyeleri:", ", ".join(f"{m.full_name} ({m.student_id})" for m in members)),
            ("Proje Başlığı:", PROJECT_TITLE),
        ]
        table = self._new_table(rows=len(info), cols=2, widths_cm=[4.0, 12.0])
        for row, (label, value) in zip(table.rows, info):
            _write_cell(row.cells[0], label, bold=True, size=BODY_FONT_SIZE)
            _write_cell(row.cells[1], value, size=BODY_FONT_SIZE)
        self._bold_line("Bölüm 2 - Ekip içi Çalışma Süreci")
        self._checklist(team_process)
        self._bold_line("Bölüm 3 - Projenin İlerleme Durumu")
        self._checklist(progress)

    def add_title_block(self) -> None:
        self.document.add_paragraph()
        for line in [*TITLE_LINES, f"{self.week}. HAFTA İLERLEME RAPORU"]:
            paragraph = self.document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(line)
            run.bold = True
            run.font.size = Pt(11)

    def add_heading(self, text: str) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(11)

    def add_box(self, paragraphs: list[str]) -> None:
        """Ders şablonundaki çerçeveli metin kutusu: tek hücreli kenarlıklı tablo."""
        cell = self._new_table(rows=1, cols=1).rows[0].cells[0]
        for index, text in enumerate(paragraphs):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(text)

    def add_labeled_box(self, items: list[tuple[str, str]]) -> None:
        """Başında altı çizili kalın etiket olan paragraflardan oluşan kutu (örn. "Amaç:", "Hedefler:")."""
        cell = self._new_table(rows=1, cols=1).rows[0].cells[0]
        for index, (label, text) in enumerate(items):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            label_run = paragraph.add_run(label)
            label_run.bold = True
            label_run.underline = True
            paragraph.add_run(f" {text}")

    def add_table(self, caption: str, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None) -> None:
        self._table_number += 1
        caption_paragraph = self.document.add_paragraph()
        caption_paragraph.paragraph_format.space_before = Pt(6)
        caption_paragraph.paragraph_format.space_after = Pt(2)
        caption_paragraph.paragraph_format.keep_with_next = True
        run = caption_paragraph.add_run(f"Tablo {self._table_number}. {caption}")
        run.bold = True
        run.font.size = TABLE_FONT_SIZE
        table = self._new_table(rows=len(rows) + 1, cols=len(headers), widths_cm=widths_cm)
        for cell, header in zip(table.rows[0].cells, headers):
            _write_cell(cell, header, bold=True)
            _shade_cell(cell, HEADER_FILL)
        for table_row, values in zip(table.rows[1:], rows):
            for cell, value in zip(table_row.cells, values):
                _write_cell(cell, value)

    def add_figure(self, image_path: Path, caption: str, width: Cm = FIGURE_WIDTH) -> None:
        if not image_path.exists():
            raise FileNotFoundError(f"Şekil bulunamadı: '{image_path}'. Önce notebooks/build_kesifsel_analiz.py çalıştırılmalı.")
        self._figure_number += 1
        picture_paragraph = self.document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.keep_with_next = True
        picture_paragraph.add_run().add_picture(str(image_path), width=width)
        caption_paragraph = self.document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(f"Şekil {self._figure_number}. {caption}")
        run.font.size = TABLE_FONT_SIZE
        run.font.color.rgb = CAPTION_COLOR

    def add_paragraph(self, text: str) -> None:
        paragraph = self.document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def add_links(self) -> None:
        self.add_heading("5. PROJE DOSYALARININ BULUNDUĞU LİNK(DRIVE, GITHUP VB.)")
        self.add_box(PROJECT_LINKS)

    def add_references(self, references: list[str]) -> None:
        self.add_heading("KAYNAKLAR")
        for number, reference in enumerate(references, start=1):
            paragraph = self.document.add_paragraph(f"{number}. {reference}")
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            paragraph.paragraph_format.space_after = Pt(3)

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(path))
        return path


__all__ = [
    "CAPTAIN_ROLE",
    "Member",
    "TEAM",
    "TEAM_PROCESS_CRITERIA",
    "WeeklyReport",
    "find_member",
    "report_file_name",
]
