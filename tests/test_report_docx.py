import json
from pathlib import Path

import pytest
from docx import Document

from reports.hafta3_rapor import (
    SUMMARY_PATH,
    build_report,
    describe_age_test,
    describe_chi_square_significance,
    describe_mann_whitney,
    describe_strong_pairs,
)
from reports.hafta3_rapor_kisa import build_short_report, short_file_name
from src.formatting import format_number, format_p_value, format_percent
from src.report_docx import TEAM, WeeklyReport, find_member, report_file_name

CAPTAIN_ID = "230260054"
MAIN_HEADINGS = [
    "TAKIM KAPTANI GERİ BİLDİRİM FORMU",
    "3. HAFTA İLERLEME RAPORU",
    "1. GİRİŞ",
    "1.1. Amaç ve Hedefler",
    "2. YÖNTEM",
    "3. PROJE YÖNETİMİ",
    "4. TEORİ VE SONUÇLAR",
    "5. PROJE DOSYALARININ BULUNDUĞU LİNK(DRIVE, GITHUP VB.)",
    "KAYNAKLAR",
]


def document_text(document) -> str:
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return "\n".join(texts)


@pytest.fixture(scope="module")
def week3_document(tmp_path_factory):
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    path = build_report(summary, CAPTAIN_ID).save(tmp_path_factory.mktemp("rapor") / "hafta3.docx")
    return Document(str(path)), summary


@pytest.mark.parametrize(
    ("ratio", "expected"),
    [(0.11698, "%11,70"), (0.6473, "%64,73"), (0.0, "%0,00")],
)
def test_format_percent_uses_turkish_decimal_comma(ratio, expected) -> None:
    assert format_percent(ratio) == expected


def test_format_number_uses_dot_for_thousands_and_comma_for_decimals() -> None:
    assert format_number(45211) == "45.211"
    assert format_number(1362.2763, 2) == "1.362,28"


def test_format_p_value_reports_tiny_values_as_threshold() -> None:
    assert format_p_value(0.0) == "< 0,001"
    assert format_p_value(0.0628) == "0,063"


def test_find_member_rejects_unknown_student_id() -> None:
    with pytest.raises(ValueError, match="ekipte yok"):
        find_member("000000000")


def test_report_file_name_follows_week_two_submission_pattern() -> None:
    assert report_file_name(3) == "230260054_230260143_230260173_hafta3.docx"


def test_captain_form_lists_captain_separately_from_members(tmp_path) -> None:
    captain = find_member("230260143")
    report = WeeklyReport(week=3, captain_id=captain.student_id)

    report.add_captain_form(team_process=[("Kriter", True)], progress=[("İlerleme", False)])
    document = Document(str(report.save(tmp_path / "form.docx")))

    info_rows = {row.cells[0].text: row.cells[1].text for row in document.tables[0].rows}
    assert info_rows["Takım Kaptanı:"] == f"{captain.full_name} ({captain.student_id})"
    assert captain.full_name not in info_rows["Ekip Üyeleri:"]
    assert all(member.full_name in info_rows["Ekip Üyeleri:"] for member in TEAM if member != captain)


def test_unchecked_progress_criterion_is_marked_in_x_column(tmp_path) -> None:
    report = WeeklyReport(week=3, captain_id=CAPTAIN_ID)

    report.add_captain_form(team_process=[("Kriter", True)], progress=[("Yapılmadı", False)])
    document = Document(str(report.save(tmp_path / "form.docx")))

    progress_row = document.tables[2].rows[1]
    assert [cell.text for cell in progress_row.cells] == ["Yapılmadı", "☐", "√"]


def test_tables_and_figures_are_numbered_in_order(tmp_path) -> None:
    figure = Path(__file__).resolve().parent.parent / "reports" / "figures" / "h3_kabul_yas.png"
    report = WeeklyReport(week=3, captain_id=CAPTAIN_ID)

    report.add_table("Birinci", ["A"], [["1"]])
    report.add_figure(figure, "Grafik")
    report.add_table("İkinci", ["A"], [["2"]])
    text = document_text(Document(str(report.save(tmp_path / "numara.docx"))))

    assert "Tablo 1. Birinci" in text
    assert "Tablo 2. İkinci" in text
    assert "Şekil 1. Grafik" in text


def test_add_figure_fails_clearly_when_image_is_missing(tmp_path) -> None:
    report = WeeklyReport(week=3, captain_id=CAPTAIN_ID)

    with pytest.raises(FileNotFoundError, match="build_kesifsel_analiz"):
        report.add_figure(tmp_path / "yok.png", "Eksik")


def test_week3_report_contains_all_main_sections(week3_document) -> None:
    document, _ = week3_document

    text = document_text(document)

    missing = [heading for heading in MAIN_HEADINGS if heading not in text]
    assert missing == []


def test_week3_report_embeds_every_figure(week3_document) -> None:
    document, _ = week3_document

    assert len(document.inline_shapes) == 8


def test_week3_report_numbers_come_from_summary(week3_document) -> None:
    document, summary = week3_document

    text = document_text(document)

    assert format_percent(summary["overall_acceptance_rate"]) in text
    assert format_percent(summary["acceptance_rates"]["poutcome"]["success"]["acceptance_rate"]) in text
    assert format_number(summary["chi_square"]["poutcome"]["cramers_v"], 3) in text
    assert format_percent(summary["age_bin_rates"]["61+"]["acceptance_rate"]) in text


def test_captain_duty_is_credited_to_the_selected_captain(tmp_path) -> None:
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    captain = find_member("230260143")

    document = Document(str(build_report(summary, captain.student_id).save(tmp_path / "kaptan.docx")))

    done_table = next(
        table for table in document.tables
        if len(table.columns) == 2 and table.rows[0].cells[1].text == "Yapılan Çalışma"
    )
    duties = {row.cells[0].text: row.cells[1].text for row in done_table.rows[1:]}
    assert duties[captain.full_name].startswith("Takım kaptanlığı")
    assert all("Takım kaptanlığı" not in text for name, text in duties.items() if name != captain.full_name)


def test_table_borders_precede_table_look_as_required_by_ooxml_schema(tmp_path) -> None:
    report = WeeklyReport(week=3, captain_id=CAPTAIN_ID)

    report.add_table("Sıra", ["A"], [["1"]])
    table = Document(str(report.save(tmp_path / "sira.docx"))).tables[0]

    child_tags = [child.tag.split("}")[1] for child in table._tbl.tblPr]
    assert child_tags.index("tblBorders") < child_tags.index("tblLayout") < child_tags.index("tblLook")


def test_mann_whitney_text_lists_non_significant_variables_from_data() -> None:
    mann_whitney = {
        "balance": {"p_value": 0.0, "rank_biserial": 0.2},
        "campaign": {"p_value": 0.0, "rank_biserial": -0.1},
        "age": {"p_value": 0.4, "rank_biserial": -0.01},
    }

    text = describe_mann_whitney(mann_whitney)

    assert "balance kabul edenlerde daha yüksek" in text
    assert "campaign kabul edenlerde daha düşük" in text
    assert "age (p = 0,400" in text


def test_chi_square_text_names_variables_that_are_not_significant() -> None:
    chi_square = {"poutcome": {"p_value": 0.0}, "default": {"p_value": 0.2}}

    text = describe_chi_square_significance(chi_square)

    assert text.startswith("default dışındaki")


def test_strong_pair_text_handles_absence_of_strong_correlations() -> None:
    assert "bulunmamıştır" in describe_strong_pairs([])


def test_age_note_appears_only_when_age_difference_is_not_significant() -> None:
    not_significant = {"age": {"p_value": 0.063, "median_yes": 38, "median_no": 39}}
    significant = {"age": {"p_value": 0.001, "median_yes": 38, "median_no": 45}}

    assert "(38 ve 39)" in describe_age_test(not_significant)
    assert describe_age_test(significant) == ""


@pytest.fixture(scope="module")
def short_document(tmp_path_factory):
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    path = build_short_report(summary, CAPTAIN_ID).save(tmp_path_factory.mktemp("kisa") / "hafta3_kisa.docx")
    return Document(str(path)), summary


def test_short_report_keeps_sections_figures_and_summary_numbers(short_document) -> None:
    document, summary = short_document

    text = document_text(document)

    assert [heading for heading in MAIN_HEADINGS if heading not in text] == []
    assert len(document.inline_shapes) == 8
    assert format_percent(summary["acceptance_rates"]["poutcome"]["success"]["acceptance_rate"]) in text
    assert format_number(summary["chi_square"]["poutcome"]["cramers_v"], 3) in text


def test_short_report_is_shorter_than_full_report(week3_document, short_document) -> None:
    full_words = len(document_text(week3_document[0]).split())
    short_words = len(document_text(short_document[0]).split())

    assert short_words < full_words


def test_short_report_lists_only_cited_references(short_document) -> None:
    text = document_text(short_document[0])

    assert all(author not in text for author in ["Wirth", "Cochran", "Kerby"])
    assert all(author in text for author in ["Spearman, C.", "Cramér, H.", "Mann, H. B."])


def test_short_report_file_name_is_marked_as_short() -> None:
    assert short_file_name() == "230260054_230260143_230260173_hafta3_kisa.docx"
